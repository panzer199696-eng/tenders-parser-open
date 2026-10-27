"""
Ссылки, даты и порядок по сроку (ревью данных 16.10.2026, Г3).

• Сортировка по сроку шла строкой «ДД.ММ.ГГГГ»: «05.11» вставало раньше
  «20.10», и ближний срок оказывался ниже дальнего.
• Лента вела на собственную ссылку лота, а не на первоисточник; у лота
  без ссылки стоял href="#" — клик открывал ту же страницу.
• В Excel срок и дата публикации были текстом: фильтр и сортировка по
  дате в книге не работали.
"""

import sys
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.core.excel_report import build_report
from src.db import connect, ensure_schema, load_feed

NOW = datetime(2026, 10, 16, 10, 0)
NOMENCLATURE = {
    "groups": [
        {"id": "reagents", "title": "Материалы цеха", "weight": 30, "tier": "daily"}
    ]
}


@pytest.fixture
def db(tmp_path):
    conn = connect(tmp_path / "t.db")
    ensure_schema(conn)
    return conn


def _add(conn, lot_id, sources=(), **over):
    row = {
        "lot_key": f"Росэлторг|{lot_id}",
        "platform": "Росэлторг",
        "lot_id": lot_id,
        "name": f"Лот {lot_id}",
        "company": "АО Рудник",
        "href": None,
        "date_published": "01.10.2026",
        "date_deadline": "30.10.2026 10:00",
        "score": 30,
        "tags": '["reagents"]',
        "core_profile": 1,
        "client_level": "market",
    }
    row.update(over)
    cols = ", ".join(row)
    marks = ", ".join("?" for _ in row)
    conn.execute(f"INSERT INTO lots ({cols}) VALUES ({marks})", list(row.values()))
    for platform, href in sources:
        conn.execute(
            "INSERT INTO lot_sources (lot_key, platform, href) VALUES (?, ?, ?)",
            (row["lot_key"], platform, href),
        )
    conn.commit()


def _book(conn):
    data = build_report(conn, today=NOW.date(), now=NOW, nomenclature=NOMENCLATURE)
    return load_workbook(BytesIO(data))


def _column(ws, title):
    head = [c.value for c in ws[1]]
    col = head.index(title) + 1
    return [ws.cell(row=r, column=col) for r in range(2, ws.max_row + 1)]


# ── Порядок по сроку ─────────────────────────────────────────────────────────


def test_лента_при_равной_оценке_ближний_срок_выше(db):
    _add(db, "1", date_deadline="05.11.2026 10:00")
    _add(db, "2", date_deadline="20.10.2026 10:00")
    _add(db, "3", date_deadline="Не указана")
    feed = load_feed(db, today=NOW.date(), now=NOW)
    assert [lot["lot_id"] for lot in feed["buckets"]["core"]] == ["2", "1", "3"]


def test_реестр_ближний_срок_выше(db):
    from src.registry import Registry

    _add(db, "1", date_deadline="05.11.2026 10:00")
    _add(db, "2", date_deadline="20.10.2026 10:00")
    reg = Registry(db)
    reg.take("Росэлторг|1", "Пётр")
    reg.take("Росэлторг|2", "Пётр")
    assert [r["lot_key"] for r in reg.list()] == ["Росэлторг|2", "Росэлторг|1"]


# ── Ссылки ───────────────────────────────────────────────────────────────────


def test_лента_ведёт_на_первоисточник(db):
    _add(
        db,
        "1",
        sources=[
            ("Bidzaar", "https://bidzaar.com/t/1"),
            ("Росэлторг", "https://www.roseltorg.ru/procedure/1"),
        ],
        href="https://bidzaar.com/t/1",
    )
    lot = load_feed(db, today=NOW.date(), now=NOW)["buckets"]["core"][0]
    assert lot["main_href"] == "https://www.roseltorg.ru/procedure/1"


def test_источник_без_ссылки_не_выбирается_для_перехода(db):
    _add(
        db,
        "1",
        sources=[("Росэлторг", None), ("Bidzaar", "https://bidzaar.com/t/1")],
    )
    lot = load_feed(db, today=NOW.date(), now=NOW)["buckets"]["core"][0]
    assert lot["main_href"] == "https://bidzaar.com/t/1"


def test_лот_без_ссылки_ищется_по_номеру(db):
    _add(
        db,
        "1",
        sources=[("ТЭК-Торг", None)],
        platform="ТЭК-Торг",
        lot_key="ТЭК-Торг|1",
        tender_no="ТЭК-555",
    )
    lot = load_feed(db, today=NOW.date(), now=NOW)["buckets"]["core"][0]
    assert not lot["main_href"]
    cell = _column(_book(db)["Все активные"], "Ссылка")[0]
    assert cell.value == "№ ТЭК-555 — искать на площадке"
    assert cell.hyperlink is None


def test_excel_ссылка_на_первоисточник(db):
    _add(
        db,
        "1",
        sources=[
            ("Bidzaar", "https://bidzaar.com/t/1"),
            ("Росэлторг", "https://www.roseltorg.ru/procedure/1"),
        ],
        href="https://bidzaar.com/t/1",
    )
    cell = _column(_book(db)["Все активные"], "Ссылка")[0]
    assert cell.hyperlink.target == "https://www.roseltorg.ru/procedure/1"


# ── Даты в Excel ─────────────────────────────────────────────────────────────


def test_excel_срок_и_публикация_ячейки_даты(db):
    _add(db, "1", date_deadline="20.10.2026 10:00", date_published="01.10.2026")
    _add(db, "2", date_deadline="Не указана", date_published="Не указана")
    ws = _book(db)["Все активные"]
    deadline = {c.row: c for c in _column(ws, "Дедлайн")}
    published = {c.row: c for c in _column(ws, "Опубликован")}
    names = {c.row: c.value for c in _column(ws, "Название")}
    row1 = next(r for r, n in names.items() if n == "Лот 1")
    row2 = next(r for r, n in names.items() if n == "Лот 2")
    assert deadline[row1].value == datetime(2026, 10, 20, 10, 0)
    assert deadline[row1].number_format == "DD.MM.YYYY HH:MM"
    assert published[row1].value.date() == date(2026, 10, 1)
    assert published[row1].number_format == "DD.MM.YYYY"
    # Непонятое остаётся текстом — исходное значение не теряем.
    assert deadline[row2].value == "Не указана"


def test_excel_лист_по_сроку_ближний_выше(db, monkeypatch):
    from src import settings_store

    monkeypatch.setattr(
        settings_store, "load_settings", lambda: {"watch_clients": ["Кедр"]}
    )
    _add(db, "1", date_deadline="05.11.2026 10:00", company='АО "КЕДР"')
    _add(db, "2", date_deadline="20.10.2026 10:00", company='АО "КЕДР"')
    names = [c.value for c in _column(_book(db)["Кедр"], "Название")]
    assert names == ["Лот 2", "Лот 1"]
