"""
«Сейчас» по Москве во всех проверках срока (ревью данных 16.10.2026, Г2).

Сроки площадок записаны по Москве, а ПК живёт в UTC+7. Ночной сбор в 03:00
по ПК — это 23:00 МСК предыдущего дня: с часами ПК лоты «до 23:59 сегодня
по Москве» отсекались как истёкшие ещё до того, как их кто-то увидел.

Здесь подменяется только московское «сейчас» в модулях; если какое-то
место по-прежнему смотрит на часы ПК, тест это покажет.
"""

import sys
from datetime import datetime
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.db import connect, ensure_schema

# 03:00 17.10 по ПК (UTC+7) = 23:00 16.10 по Москве.
MSK_NOW = datetime(2026, 10, 16, 23, 0)


@pytest.fixture
def msk(monkeypatch):
    """Подменить московское «сейчас» везде, где его спрашивают."""
    from src import dates
    from src.core import excel_report
    from src.db import feed, items, reports
    from src.profile import profile
    from src.registry import registry
    from src.utils import classifier

    # src.dates — тоже: resolve_now спрашивает «сейчас» у него самого.
    modules = (dates, feed, items, reports, profile, registry, classifier, excel_report)
    for module in modules:
        monkeypatch.setattr(module, "now_msk", lambda: MSK_NOW, raising=False)
        monkeypatch.setattr(module, "today_msk", lambda: MSK_NOW.date(), raising=False)
    return MSK_NOW


@pytest.fixture
def db(tmp_path):
    conn = connect(tmp_path / "t.db")
    ensure_schema(conn)
    return conn


def _add(conn, lot_id, deadline, **over):
    row = {
        "lot_key": f"Росэлторг|{lot_id}",
        "platform": "Росэлторг",
        "lot_id": lot_id,
        "name": "Компаунд заливочный",
        "company": "АО Рудник",
        "href": f"https://www.roseltorg.ru/app/market-next/x/tender-{lot_id}/",
        "date_deadline": deadline,
        "score": 50,
        "tags": '["reagents"]',
        "client_level": "market",
        "core_profile": 1,
        "match_strength": "strong",
    }
    row.update(over)
    cols = ", ".join(row)
    marks = ", ".join("?" for _ in row)
    conn.execute(f"INSERT INTO lots ({cols}) VALUES ({marks})", list(row.values()))
    conn.commit()


def test_лента_сегодняшний_по_москве_лот_не_истёк(db, msk):
    from src.db import load_feed

    _add(db, "1", "16.10.2026 23:30")
    _add(db, "2", "16.10.2026 22:00")
    feed = load_feed(db)
    assert [lot["lot_id"] for lot in feed["buckets"]["urgent"]] == ["1"]
    assert [lot["lot_id"] for lot in feed["buckets"]["expired"]] == ["2"]
    assert feed["buckets"]["urgent"][0]["days_left"] == 0


def test_дни_до_срока_по_московской_дате(msk):
    from src.db.feed import days_left

    assert days_left("17.10.2026 10:00") == 1


def test_фильтр_активности_не_режет_сегодняшние_по_москве(msk):
    from src.utils.classifier import is_active

    assert is_active({"date_deadline": "16.10.2026 23:59"})
    assert not is_active({"date_deadline": "15.10.2026 23:59"})


def test_профиль_истечение_по_москве(msk):
    from src.profile.profile import Profile

    assert not Profile.is_expired("16.10.2026 23:30")
    assert Profile.is_expired("16.10.2026 22:00")
    assert Profile.days_until("17.10.2026") == 1


def test_отчёт_сроки_считает_от_текущего_момента(db, msk):
    """Раньше «сейчас» в отчёте было полночью: лот, истёкший сегодня
    утром, весь день числился горящим."""
    from src.db.reports import deadlines

    _add(db, "1", "16.10.2026 09:00")
    _add(db, "2", "16.10.2026 23:30")
    result = deadlines(db)
    assert result["expired"] == 1
    assert result["urgent"] == 1
    assert result["as_of"] == "2026-10-16"


def test_отчёт_сроки_на_заданный_момент(db):
    from src.db.reports import deadlines

    _add(db, "1", "01.09.2026 09:00")
    result = deadlines(db, now=datetime(2026, 9, 1, 10, 0))
    assert result["expired"] == 1
    assert result["urgent"] == 0
    assert result["as_of"] == "2026-09-01"


def test_красная_зона_по_москве(db, msk):
    from src.registry import Registry

    _add(db, "1", "16.10.2026 23:30")
    reg = Registry(db)
    reg.take("Росэлторг|1", "Пётр")
    zone = reg.red_zone()
    assert [r["lot_key"] for r in zone["горит"]] == ["Росэлторг|1"]
    assert zone["просрочено"] == []


def test_excel_сегодняшний_по_москве_в_книге(db, msk):
    from io import BytesIO

    from openpyxl import load_workbook

    from src.core.excel_report import build_report

    _add(db, "1", "16.10.2026 23:30")
    _add(db, "2", "16.10.2026 22:00")
    wb = load_workbook(BytesIO(build_report(db, nomenclature={"groups": []})))
    ws = wb["Все активные"]
    head = [c.value for c in ws[1]]
    days = head.index("Осталось дней") + 1
    assert ws.max_row == 2, "в книге только неистёкший по Москве лот"
    assert ws.cell(2, days).value == 0


# ── Одна пара «сегодня/сейчас» на всю программу ──────────────────────────────


def test_сегодня_берётся_из_заданного_момента():
    from src.dates import resolve_now

    at = datetime(2026, 6, 1, 10, 0)
    assert resolve_now(now=at) == (at.date(), at)
    day, now = resolve_now(today=at.date())
    assert day == at.date() and now.date() == at.date()


def test_книга_на_заданный_момент_считает_дни_от_него(db):
    """Задан только момент — день берётся из него, а не с часов ПК: иначе
    «Осталось дней» в книге на 01.06 считалось от сегодняшнего числа."""
    from io import BytesIO

    from openpyxl import load_workbook

    from src.core.excel_report import build_report

    _add(db, "1", "03.06.2026 10:00")
    data = build_report(
        db, now=datetime(2026, 6, 1, 10, 0), nomenclature={"groups": []}
    )
    ws = load_workbook(BytesIO(data))["Все активные"]
    head = [c.value for c in ws[1]]
    assert ws.cell(2, head.index("Осталось дней") + 1).value == 2


def test_лента_на_заданный_момент_считает_дни_от_него(db):
    from src.db import load_feed

    _add(db, "1", "03.06.2026 10:00")
    feed = load_feed(db, now=datetime(2026, 6, 1, 10, 0))
    assert feed["buckets"]["urgent"][0]["days_left"] == 2
