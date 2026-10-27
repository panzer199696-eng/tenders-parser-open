"""
Одна раскладка по корзинам для ленты, Excel и отчётов (ревью данных, Г4).

До 16.10.2026 раскладок было три: лента («горит» — только наше), Excel
(«горит» — любой лот с оценкой > 0, взятые и отклонённые тоже в книге) и
отчёт «Сроки» (своё правило и полночь вместо «сейчас»). Одна база давала
три разных ответа на вопрос «что горит». Решение владельца: «Горит» — по
правилу ленты, взятые и «Мимо» в книге не показывать.

Главный тест здесь — инвариант: на одной базе корзина ленты, лист книги
и цифра отчёта совпадают.
"""

import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.core.excel_report import build_report
from src.db import connect, ensure_schema, load_feed
from src.db.feed import classify_row
from src.db.reports import deadlines

NOW = datetime(2026, 10, 16, 10, 0)
TODAY = NOW.date()

NOMENCLATURE = {
    "groups": [
        {"id": "reagents", "title": "Материалы цеха", "weight": 30, "tier": "daily"},
        {"id": "pumps", "title": "Насосы", "weight": 20, "tier": "off"},
    ]
}


@pytest.fixture
def db(tmp_path, monkeypatch):
    from src import settings_store

    monkeypatch.setattr(
        settings_store, "load_settings", lambda: {"watch_clients": ["Кедр"]}
    )
    conn = connect(tmp_path / "t.db")
    ensure_schema(conn)
    return conn


def _add(conn, lot_id, name, **over):
    row = {
        "lot_key": f"Росэлторг|{lot_id}",
        "platform": "Росэлторг",
        "lot_id": lot_id,
        "name": name,
        "company": "АО Рудник",
        "href": f"https://www.example.invalid/lot/{lot_id}",
        "date_deadline": "30.10.2026 10:00",
        "score": 30,
        "tags": '["reagents"]',
        "client_level": "market",
        "industry": "ours",
        "match_strength": "strong",
        "core_profile": 1,
        "first_seen_at": "2026-10-01T03:10:00",
        "last_seen_at": "2026-10-16T03:10:00",
    }
    row.update(over)
    cols = ", ".join(row)
    marks = ", ".join("?" for _ in row)
    conn.execute(f"INSERT INTO lots ({cols}) VALUES ({marks})", list(row.values()))
    conn.execute(
        "INSERT INTO lot_sources (lot_key, platform, href, lot_id, seen_at)"
        " VALUES (?, ?, ?, ?, ?)",
        (row["lot_key"], row["platform"], row["href"], lot_id, row["last_seen_at"]),
    )
    conn.commit()


def _fill(conn):
    _add(conn, "1", "Горит профиль", date_deadline="18.10.2026 10:00")
    _add(conn, "2", "Профиль")
    _add(conn, "3", "Клиент без темы", tags="[]", core_profile=0, client_level="full")
    _add(conn, "4", "Ниже порога", below_threshold=1)
    _add(
        conn,
        "5",
        "Слабое",
        tags='["pumps"]',
        core_profile=0,
        needs_review=1,
        match_strength="weak",
    )
    _add(conn, "6", "Чужая отрасль", off_industry=1, industry="foreign")
    _add(
        conn,
        "7",
        "Кедр подряд",
        off_profile=1,
        core_profile=0,
        company='АО "ЗРК "КЕДР"',
    )
    _add(conn, "8", "Истёк", date_deadline="15.10.2026 10:00")
    _add(conn, "9", "Взят в работу", date_deadline="17.10.2026 10:00")
    conn.execute(
        "INSERT INTO applications (lot_key, status) VALUES ('Росэлторг|9', 'SELECTED')"
    )
    _add(conn, "10", "Отклонён")
    conn.execute(
        "INSERT INTO applications (lot_key, status) VALUES ('Росэлторг|10', 'REJECTED')"
    )
    # Рынок, тема угадана, срок завтра: по правилу ленты не «горит».
    _add(
        conn,
        "11",
        "Рынок насосы завтра",
        tags='["pumps"]',
        core_profile=0,
        date_deadline="17.10.2026 10:00",
        score=20,
    )
    _add(
        conn,
        "12",
        "Новый горящий",
        date_deadline="17.10.2026 12:00",
        first_seen_at="2026-10-16T03:20:00",
    )
    # Площадка не показывает лот с 10.10 — две полные ночи подряд.
    _add(conn, "13", "Пропал", last_seen_at="2026-10-10T03:10:00")
    for day in ("14", "15", "16"):
        cur = conn.execute(
            "INSERT INTO runs (started_at, finished_at, outcome) VALUES (?, ?, 'ok')",
            (f"2026-10-{day}T03:00:00", f"2026-10-{day}T03:40:00"),
        )
        conn.execute(
            "INSERT INTO run_platforms (run_id, platform, complete, lots)"
            " VALUES (?, 'Росэлторг', 1, 100)",
            (cur.lastrowid,),
        )
    conn.commit()


def _book(conn):
    data = build_report(conn, today=TODAY, now=NOW, nomenclature=NOMENCLATURE)
    return load_workbook(BytesIO(data))


def _sheet_names(ws):
    head = [c.value for c in ws[1]]
    col = head.index("Название") + 1
    return sorted(ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1))


def _feed_names(feed, key):
    return sorted(lot["name"] for lot in feed["buckets"][key])


SHEETS = {
    "urgent": "🔥 Горит",
    "client": "Закреплённые клиенты",
    "below": "Ниже порога",
    "review": "На проверку",
    "foreign": "Не наш профиль",
    "missing": "Пропал — проверить",
}


def test_корзина_ленты_равна_листу_книги_и_отчёту(db):
    _fill(db)
    feed = load_feed(db, today=TODAY, now=NOW)
    wb = _book(db)
    for key, sheet in SHEETS.items():
        assert _sheet_names(wb[sheet]) == _feed_names(feed, key), sheet
    active = sorted(
        lot["name"]
        for key, lots in feed["buckets"].items()
        if key != "expired"
        for lot in lots
    )
    assert _sheet_names(wb["Все активные"]) == active
    report = deadlines(db, now=NOW)
    assert report["urgent"] == feed["counts"]["urgent"]
    assert report["missing"] == feed["counts"]["missing"]


def test_горит_по_правилу_ленты(db):
    _fill(db)
    names = _sheet_names(_book(db)["🔥 Горит"])
    assert names == ["Горит профиль", "Новый горящий"]
    assert "Рынок насосы завтра" not in names


def test_уверенный_второй_эшелон_виден_на_листе_темы(db):
    """Лента копит уверенно распознанный рынок в «На проверку»; в книге его
    ищут по листу темы — лист темы его не теряет."""
    _fill(db)
    wb = _book(db)
    assert _sheet_names(wb["Насосы"]) == ["Рынок насосы завтра"]
    assert "Рынок насосы завтра" in _sheet_names(wb["На проверку"])
    assert "Слабое" not in _sheet_names(wb["Насосы"])


def test_взятые_и_отклонённые_в_книгу_не_попадают(db):
    _fill(db)
    names = _sheet_names(_book(db)["Все активные"])
    assert "Взят в работу" not in names
    assert "Отклонён" not in names


def test_особый_контроль_отсеянный_профилем_в_закреплённых(db):
    _fill(db)
    wb = _book(db)
    assert "Кедр подряд" in _sheet_names(wb["Закреплённые клиенты"])
    assert "Кедр подряд" not in _sheet_names(wb["Не наш профиль"])
    feed = load_feed(db, today=TODAY, now=NOW)
    assert "Кедр подряд" in _feed_names(feed, "client")


def test_свежие_горящие_считаются_по_корзине(db):
    _fill(db)
    fresh = load_feed(db, today=TODAY, now=NOW)["fresh"]
    assert fresh["total"] == 1
    assert fresh["urgent"] == 1


def test_classify_row_одна_функция_раскладки():
    row = {
        "date_deadline": "17.10.2026 10:00",
        "core_profile": 0,
        "client_level": "market",
        "tags_list": ["pumps"],
        "company": "АО Рудник",
    }
    assert classify_row(row, NOW, watched=[]) == "review"
    assert classify_row({**row, "core_profile": 1}, NOW, watched=[]) == "urgent"
    assert classify_row({**row, "missing": True}, NOW, watched=[]) == "missing"
    expired = {**row, "date_deadline": "15.10.2026 10:00", "missing": True}
    assert classify_row(expired, NOW, watched=[]) == "expired"
