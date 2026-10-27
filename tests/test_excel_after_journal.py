"""Книга Excel строится ПОСЛЕ записи прогона (ревью 16.10.2026, В2).

Инцидент: main_ui строил книгу до записи журнала, и шапка «Когда собрано»
с блоком полноты брались из ПРЕДЫДУЩЕГО прогона — в ночном файле 16.10
стояло «15.10 в 07:06». Теперь запись прогона идёт первой, книга получает
его run_id; для догоняющего прогона блок полноты — ночной, дополненный
догнанными площадками. Сбой Excel не делает прогон ошибкой.
"""

from datetime import date, datetime
from io import BytesIO
from unittest.mock import AsyncMock, Mock

import pytest
from openpyxl import load_workbook

from src.core import collect
from src.core import context as ctx
from src import run_journal
from src.core.excel_report import build_report
from src.db import connect, ensure_schema
from src.db.runs import record_run, run_platforms_for
from src.parsers import completeness as c

NOMENCLATURE = {"groups": []}


def _rec(started, finished, rows, trigger=None, outcome="ok"):
    return {
        "started_at": started,
        "finished_at": finished,
        "outcome": outcome,
        "trigger": trigger,
        "data_complete": all(r.get("data_complete") for r in rows),
        "by_platform": rows,
        "platforms": {},
    }


def _row(platform, complete, reason=""):
    return {
        "platform": platform,
        "units": 3,
        "data_complete": complete,
        "reason": reason,
        "verdict": "собрано целиком" if complete else reason,
    }


@pytest.fixture
def conn(tmp_path):
    db = connect(tmp_path / "t.db")
    ensure_schema(db)
    yield db
    db.close()


def _sheet(conn, **kw):
    data = build_report(
        conn,
        today=date(2026, 10, 16),
        now=datetime(2026, 10, 16, 8),
        nomenclature=NOMENCLATURE,
        **kw,
    )
    return load_workbook(BytesIO(data)).worksheets[0]


def _block(ws):
    """Строки блока полноты: площадка → вердикт."""
    out = {}
    for row in ws.iter_rows(min_row=4, max_col=3, values_only=True):
        if row[0] and str(row[0]).startswith("   "):
            out[row[0].strip()] = row[2]
    return out


def test_шапка_книги_берётся_из_переданного_прогона(conn):
    old = record_run(
        conn,
        _rec(
            "2026-10-15T07:00:01",
            "2026-10-15T07:06:58",
            [_row("ТЭК-Торг", True)],
            trigger="catchup",
        ),
    )
    new = record_run(
        conn,
        _rec("2026-10-16T03:00:01", "2026-10-16T03:32:20", [_row("Bidzaar", True)]),
    )
    assert _sheet(conn, run_id=new)["A1"].value == "Когда собрано: 16.10 в 03:32"
    assert "15.10 в 07:06" in _sheet(conn, run_id=old)["A1"].value


def test_догон_дополняет_ночной_блок_полноты(conn):
    night = _rec(
        "2026-10-15T03:00:01",
        "2026-10-15T03:44:42",
        [_row("Росэлторг", True), _row("ТЭК-Торг", False, "обрыв связи с площадкой")],
    )
    record_run(conn, night)
    catchup = record_run(
        conn,
        _rec(
            "2026-10-15T07:00:01",
            "2026-10-15T07:06:58",
            [_row("ТЭК-Торг", True)],
            trigger="catchup",
        ),
    )
    rows = {r["platform"]: r for r in run_platforms_for(conn, catchup)}
    assert set(rows) == {"Росэлторг", "ТЭК-Торг"}
    assert rows["ТЭК-Торг"]["complete"] == 1
    assert "догнано в 07:06" in rows["ТЭК-Торг"]["reason"]
    assert rows["Росэлторг"]["complete"] == 1

    block = _block(_sheet(conn, run_id=catchup))
    assert "догнано" in block["ТЭК-Торг"]
    assert block["Росэлторг"] == "собрано целиком"


def test_без_run_id_книга_берёт_последний_прогон(conn):
    record_run(
        conn,
        _rec("2026-10-15T03:00:01", "2026-10-15T03:44:42", [_row("Bidzaar", True)]),
    )
    record_run(
        conn,
        _rec(
            "2026-10-16T03:00:01",
            "2026-10-16T03:32:20",
            [_row("Bidzaar", False, "обрыв связи с площадкой")],
        ),
    )
    ws = _sheet(conn)
    assert ws["A1"].value == "Когда собрано: 16.10 в 03:32"
    assert _block(ws)["Bidzaar"] == "обрыв связи с площадкой"


# ── main_ui: порядок «журнал → книга» и сбой Excel ─────────────────────────


@pytest.fixture
def сбор(monkeypatch, tmp_path):
    import src.core.parser_runner as pr
    import src.keywords_store as ks
    import src.settings_store as ss
    import src.utils.browser_manager as bm
    from src import profile

    monkeypatch.setattr(collect, "_parse_state",
        {
            **collect._parse_state,
            "running": False,
            "stopping": False,
            "stopped": False,
            "error": None,
            "progress": [],
            "tenders": [],
            "excel_path": None,
            "data_saved": False,
        },
    )
    monkeypatch.setattr(collect, "_parse_cancel", {"loop": None, "task": None})
    monkeypatch.setattr(ctx, "UNIFIED_DB_PATH", tmp_path / "unified.db")
    # книга — во временную папку, не в «Документы» пользователя
    monkeypatch.setattr(collect, "results_dir", lambda: tmp_path)
    monkeypatch.setattr(run_journal, "journal_dir", lambda: tmp_path / "journal")
    monkeypatch.setattr(ctx, "load_companies", list)
    monkeypatch.setattr(ks, "enabled_phrases", list)
    monkeypatch.setattr(ss, "load_settings", lambda: {"parse_bidzaar": True})
    monkeypatch.setattr(profile, "Profile", lambda: None)
    browser = Mock(start=AsyncMock(), close=AsyncMock())
    monkeypatch.setattr(bm, "BrowserManager", Mock(return_value=browser))
    runner = Mock(
        unit_reports=[
            c.build_unit_report(
                "phrase",
                "Герметик",
                platform="Bidzaar",
                expected_pages=1,
                collected_pages=1,
            )
        ]
    )
    runner.platform_lots = {"Bidzaar": 1}
    runner.platform_newest = {}
    runner.run_all = AsyncMock(
        return_value=[
            {
                "platform": "Bidzaar",
                "lot_id": "z-1",
                "href": "https://bz.invalid/1",
                "name": "Герметик",
            }
        ]
    )
    monkeypatch.setattr(pr, "ParserRunner", Mock(return_value=runner))

    # Прошлый прогон в базе — его время книга показывать НЕ должна.
    db = connect(tmp_path / "unified.db")
    ensure_schema(db)
    record_run(
        db,
        _rec(
            "2026-10-15T07:00:01",
            "2026-10-15T07:06:58",
            [_row("ТЭК-Торг", True)],
            trigger="catchup",
        ),
    )
    db.close()
    return tmp_path


def test_A1_книги_это_только_что_записанный_прогон(сбор):
    assert collect._run_collect_once() == collect.EXIT_OK
    path = collect._parse_state["excel_path"]
    assert path
    ws = load_workbook(path).worksheets[0]
    record = run_journal.read_last_run()
    finished = datetime.fromisoformat(record["finished_at"])
    assert ws["A1"].value == f"Когда собрано: {finished:%d.%m} в {finished:%H:%M}"
    assert "15.10 в 07:06" not in ws["A1"].value
    assert set(_block(ws)) == {"Bidzaar"}


def test_сбой_excel_не_делает_прогон_ошибкой(сбор, monkeypatch):
    monkeypatch.setattr(collect, "_save_excel_file", Mock(side_effect=PermissionError("файл открыт в Excel"))
    )
    assert collect._run_collect_once() == collect.EXIT_OK
    assert collect._parse_state["error"] is None
    assert collect._parse_state["excel_path"] is None
    assert "Excel" in collect._parse_state["warning"]
    assert run_journal.read_last_run()["outcome"] == run_journal.OUTCOME_OK
    assert any("✅ Готово" in s and "Excel" in s for s in collect._parse_state["progress"])
