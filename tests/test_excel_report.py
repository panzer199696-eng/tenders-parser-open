"""
Тесты Excel-отчёта по темам из единой базы (ТЗ-5, задача 3).

Что здесь проверяется в первую очередь: **отсеянное не должно попадать на
профильные листы**. Именно с этой жалобы началась работа — в прежней книге
лист с реагентами открывался строкой «Установка хранилища серной кислоты.
Выполнение строительно-монтажных работ», потому что экспорт ничего не знал
про решения профиля.

Сеть и папка пользователя не трогаются: справочник подкладывается прямо в
вызов, база живёт в tmp_path.
"""

import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.core.excel_report import build_report
from src.db import connect, ensure_schema

TODAY = date(2026, 9, 1)
NOW = datetime(2026, 9, 1, 10, 0)

NOMENCLATURE = {
    "groups": [
        {
            "id": "feeders_own",
            "title": "Дозаторы — своё производство",
            "weight": 40,
            "enabled": True,
            "tier": "daily",
        },
        {
            "id": "reagents_zif",
            "title": "Материалы цеха",
            "weight": 30,
            "enabled": True,
            "tier": "daily",
        },
        {
            "id": "pumps_slurry",
            "title": "Насосы и насосное оборудование",
            "weight": 25,
            "enabled": True,
            "tier": "off",
        },
        {
            "id": "construction",
            "title": "Строительство и подряд",
            "weight": 5,
            "enabled": False,
            "tier": "off",
        },
    ]
}


@pytest.fixture
def db(tmp_path):
    conn = connect(tmp_path / "t.db")
    ensure_schema(conn)
    return conn


def _add(conn, lot_id, **over):
    row = {
        "lot_key": f"EIS|{lot_id}",
        "platform": "ЕИС",
        "lot_id": lot_id,
        "name": "Компаунд заливочный 20 т",
        "company": "АО Рудник",
        "href": f"https://example.test/tender-{lot_id}/",
        "date_published": "25.08.2026",
        "date_deadline": "20.09.2026",
        "score": 50,
        "score_explain": "Материалы цеха +30 · совпало точно: «компаунд заливочный»",
        "tags": '["reagents_zif"]',
        "client_level": "market",
        "industry": "ours",
        "match_strength": "strong",
        "core_profile": 1,
        "needs_review": 0,
        "off_industry": 0,
        "off_profile": 0,
        "prepayment": "unknown",
    }
    row.update(over)
    cols = ", ".join(row)
    marks = ", ".join("?" for _ in row)
    conn.execute(f"INSERT INTO lots ({cols}) VALUES ({marks})", list(row.values()))
    conn.commit()


def _book(conn):
    from io import BytesIO

    data = build_report(conn, today=TODAY, now=NOW, nomenclature=NOMENCLATURE)
    return load_workbook(BytesIO(data))


def _names(ws, column="Название"):
    head = [c.value for c in ws[1]]
    col = head.index(column) + 1
    return [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]


# ── Состав книги ─────────────────────────────────────────────────────────────


def test_книга_собирается_на_пустой_базе(db):
    """Пустая база — это книга с пустыми листами, а не отказ.

    Прежний экспорт отвечал «Нет данных для экспорта, сначала запустите
    парсинг»; человек при этом видел ленту с лотами и не понимал, почему
    выгрузка невозможна.
    """
    wb = _book(db)
    assert "Сводка по темам" in wb.sheetnames
    assert "Не наш профиль" in wb.sheetnames


def test_лист_заводится_на_каждую_непустую_тему(db):
    _add(db, "1", tags='["reagents_zif"]')
    _add(db, "2", tags='["feeders_own"]', name="Дозатор плунжерный НШ 2-12")
    wb = _book(db)
    assert "Материалы цеха" in wb.sheetnames
    assert "Дозаторы — своё производство" in wb.sheetnames
    # Насосов в базе нет — пустого листа быть не должно.
    assert "Насосы и насосное оборудование" not in wb.sheetnames


def test_темы_идут_по_весу_группы(db):
    _add(db, "1", tags='["reagents_zif"]')
    _add(db, "2", tags='["feeders_own"]')
    wb = _book(db)
    order = [
        n
        for n in wb.sheetnames
        if n in ("Дозаторы — своё производство", "Материалы цеха")
    ]
    assert order == ["Дозаторы — своё производство", "Материалы цеха"], (
        "тема с большим весом должна идти выше"
    )


def test_выключенная_группа_листа_не_получает(db):
    _add(db, "1", tags='["construction"]', name="Строительство ЛЭП")
    wb = _book(db)
    assert "Строительство и подряд" not in wb.sheetnames


def test_сводка_показывает_и_пустые_темы(db):
    """Пустая тема — сигнал пересмотреть фразы, а не повод спрятать строку."""
    _add(db, "1", tags='["reagents_zif"]')
    ws = _book(db)["Сводка по темам"]
    listed = {ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)}
    assert "Насосы и насосное оборудование" in listed
    assert "Дозаторы — своё производство" in listed


# ── Главное: отсеянное не попадает к профильному ─────────────────────────────


def test_подряд_не_попадает_на_лист_темы(db):
    _add(
        db,
        "1",
        name="Хранилище серной кислоты. Выполнение строительно-монтажных работ",
        off_profile=1,
        score=20,
        score_explain="не наш профиль: это подряд, а не поставка",
    )
    _add(db, "2", name="Компаунд заливочный 20 т")
    wb = _book(db)
    assert _names(wb["Материалы цеха"]) == ["Компаунд заливочный 20 т"]
    assert len(_names(wb["Не наш профиль"])) == 1


def test_чужая_отрасль_не_попадает_на_лист_темы(db):
    _add(
        db,
        "1",
        name="Реагент гипохлорит натрия для Сормовской ТЭЦ",
        company="ПАО Пример-Энерго",
        off_industry=1,
        industry="foreign",
        score=0,
    )
    _add(db, "2", name="Компаунд заливочный 20 т")
    wb = _book(db)
    assert _names(wb["Материалы цеха"]) == ["Компаунд заливочный 20 т"]
    assert _names(wb["Не наш профиль"]) == [
        "Реагент гипохлорит натрия для Сормовской ТЭЦ"
    ]


def test_отсеянное_не_горит(db):
    """Чужой лот с ближним дедлайном не должен отвлекать в «Горит»."""
    _add(
        db,
        "1",
        name="Кислота серная для ТЭЦ",
        date_deadline="02.09.2026",
        off_industry=1,
        score=30,
    )
    wb = _book(db)
    assert _names(wb["🔥 Горит"]) == []


def test_горит_только_ближний_срок(db):
    _add(db, "1", name="Герметик Пример", date_deadline="03.09.2026")
    _add(db, "2", name="Компаунд заливочный 20 т", date_deadline="25.09.2026")
    assert _names(_book(db)["🔥 Горит"]) == ["Герметик Пример"]


def test_истёкший_лот_в_книгу_не_попадает(db):
    _add(db, "1", name="Вчерашний лот", date_deadline="31.08.2026")
    _add(db, "2", name="Компаунд заливочный 20 т")
    names = _names(_book(db)["Все активные"])
    assert names == ["Компаунд заливочный 20 т"]


def test_слабое_совпадение_уходит_на_проверку(db):
    _add(
        db,
        "1",
        name="Насосное оборудование",
        tags='["pumps_slurry"]',
        needs_review=1,
        match_strength="weak",
        score=9,
        core_profile=0,
    )
    wb = _book(db)
    assert _names(wb["На проверку"]) == ["Насосное оборудование"]
    assert "Насосы и насосное оборудование" not in wb.sheetnames


def test_лот_без_темы_уходит_на_проверку(db):
    _add(db, "1", name="Прочие материалы", tags="[]", core_profile=0, score=0)
    assert _names(_book(db)["На проверку"]) == ["Прочие материалы"]


# ── Содержимое строки ────────────────────────────────────────────────────────


def test_почему_берётся_из_базы_как_есть(db):
    """Объяснение не пересчитывается при выгрузке: в книге ровно то, что
    программа решила при сборе, иначе лента и файл разойдутся."""
    _add(db, "1", score_explain="Материалы цеха +30 · закреплённый клиент +40")
    assert _names(_book(db)["Материалы цеха"], "Почему") == [
        "Материалы цеха +30 · закреплённый клиент +40"
    ]


def test_предоплата_не_найдено_не_выдаётся_за_нет(db):
    _add(db, "1", prepayment="unknown")
    assert _names(_book(db)["Материалы цеха"], "Предоплата") == ["не найдено"]


def test_ссылка_кликабельна(db):
    _add(db, "1")
    ws = _book(db)["Материалы цеха"]
    head = [c.value for c in ws[1]]
    cell = ws.cell(row=2, column=head.index("Ссылка") + 1)
    assert cell.hyperlink is not None
    assert cell.hyperlink.target == "https://example.test/tender-1/"


def test_имя_листа_не_длиннее_предела_excel(db):
    long_title = "Очень длинное название темы, которое Excel не примет целиком"
    nomenclature = {
        "groups": [
            {
                "id": "long",
                "title": long_title,
                "weight": 10,
                "enabled": True,
                "tier": "off",
            }
        ]
    }
    _add(db, "1", tags='["long"]')
    from io import BytesIO

    wb = load_workbook(
        BytesIO(build_report(db, today=TODAY, now=NOW, nomenclature=nomenclature))
    )
    made = [n for n in wb.sheetnames if n.startswith("Очень длинное")]
    assert made and len(made[0]) <= 31
    assert not any(ch in made[0] for ch in "[]:*?/\\")


def test_поток_закреплённого_клиента_виден_отдельным_листом(db):
    """У заказчика, с которым работаем, смотрим весь поток — даже без темы.

    Наша позиция часто спрятана внутри сборного лота («Расходники, реактивы,
    посуда»), где ни одна фраза не сработает. Без своего листа такой лот
    растворялся бы в общей куче «На проверку».
    """
    _add(db, "1", name="Расходники, реактивы, посуда", tags="[]",
         client_level="full", core_profile=0, score=40)
    wb = _book(db)
    assert _names(wb["Закреплённые клиенты"]) == ["Расходники, реактивы, посуда"]
    assert _names(wb["На проверку"]) == [], "лот не должен дублироваться"


def test_клиентский_лист_не_дублирует_листы_тем(db):
    _add(db, "1", name="Компаунд заливочный 20 т", client_level="full")
    wb = _book(db)
    assert _names(wb["Материалы цеха"]) == ["Компаунд заливочный 20 т"]
    assert _names(wb["Закреплённые клиенты"]) == []


def test_сводка_показывает_полноту_по_площадкам(db):
    """Книга обязана говорить не только «сколько нашли», но и «всё ли смогли».

    Площадка, отдавшая половину, выглядит в отчёте так же, как отдавшая всё,
    если о полноте молчать. Поэтому строка на каждую площадку последнего
    прогона: сколько объектов обошли и собран ли каждый целиком.
    """
    db.execute(
        "INSERT INTO runs (started_at, finished_at, outcome, data_complete, summary)"
        " VALUES ('2026-09-04T03:00:00', '2026-09-04T03:35:00', 'ok', 0, 'тест')"
    )
    run_id = db.execute("SELECT id FROM runs ORDER BY id DESC LIMIT 1").fetchone()[0]
    db.executemany(
        "INSERT INTO run_platforms (run_id, platform, collected, complete, reason)"
        " VALUES (?, ?, ?, ?, ?)",
        [
            (run_id, "ЕИС", 40, 1, None),
            (run_id, "Bidzaar", 106, 0, "упёрлись в потолок страниц"),
            (run_id, "Росэлторг", 7, None, None),
        ],
    )
    db.commit()
    _add(db, "1")

    ws = _book(db)["Сводка по темам"]
    text = "\n".join(
        str(ws.cell(r, c).value or "")
        for r in range(1, 12)
        for c in range(1, 4)
    )
    assert "Полнота сбора по площадкам" in text
    assert "ЕИС" in text and "собрано целиком" in text
    assert "упёрлись в потолок страниц" in text, "виновник должен быть назван"
    assert "не сообщила объём" in text, "«не знаю» не выдаём за «всё собрано»"


def test_заказчики_особого_контроля_идут_отдельным_листом(db, monkeypatch):
    """Просьба владельца 07.09.2026: Кедр и Сосна — отдельным листом.

    По ним работа идёт адресно, и искать их закупки среди трёх тысяч строк
    незачем. Список правится в настройках, а не в коде.
    """
    import src.settings_store as settings_store

    monkeypatch.setattr(
        settings_store, "load_settings", lambda: {"watch_clients": ["Кедр", "Сосна"]}
    )
    _add(db, "1", name="Лента конвейерная", company='АО "ЗОЛОТОРУДНАЯ КОМПАНИЯ КЕДР"')
    _add(db, "2", name="Футеровки бункера Metso", company='АО "ИК "СОСНА"')
    _add(db, "3", name="Компаунд заливочный 20 т", company='АО "РОМАШКА"')

    wb = _book(db)
    assert "Кедр и Сосна" in wb.sheetnames
    names = _names(wb["Кедр и Сосна"])
    assert sorted(names) == ["Лента конвейерная", "Футеровки бункера Metso"]


def test_особый_контроль_сверяет_целые_слова(db, monkeypatch):
    """«Кедрова» — не наш Кедр.

    Подстрока здесь врёт: на лист приезжали санаторий и детский лагерь имени
    Кедрова (живая база, 07.09.2026).
    """
    import src.settings_store as settings_store

    monkeypatch.setattr(settings_store, "load_settings", lambda: {"watch_clients": ["Кедр"]})
    _add(db, "1", name="Наш лот", company='АО "ЗОЛОТОРУДНАЯ КОМПАНИЯ КЕДР"')
    _add(db, "2", name="Чужой лот",
         company='МБУ "ЗАГОРОДНЫЙ ОЗДОРОВИТЕЛЬНЫЙ КОМПЛЕКС ИМЕНИ КЕДРОВА"')

    assert _names(_book(db)["Кедр"]) == ["Наш лот"]


def test_без_списка_особого_контроля_лист_не_создаётся(db, monkeypatch):
    import src.settings_store as settings_store

    monkeypatch.setattr(settings_store, "load_settings", lambda: {"watch_clients": []})
    _add(db, "1")
    wb = _book(db)
    assert not any("Кедр" in n for n in wb.sheetnames)


# ── Состав закупки (ТЗ-10) ───────────────────────────────────────────────────


def test_колонка_состав_первые_три_позиции_и_ещё(db):
    """Состав сборного лота — главный ответ на «наше/не наше» по названию,
    которое не совпало с номенклатурой. Ячейка длинная, поэтому первые три
    наименования и «…ещё N»: строки длиннее трёх позиций лист не читается.
    """
    from src.db.items import save_lot_items

    _add(db, "1", name="Расходники, реактивы, посуда")
    _add(db, "2", name="Компаунд заливочный 20 т")
    save_lot_items(
        db,
        "EIS|1",
        [
            {"position": i + 1, "name": n}
            for i, n in enumerate(
                ["Перчатки", "Бутыли", "Стаканы", "Пипетки", "Штативы"]
            )
        ],
    )

    ws = _book(db)["Материалы цеха"]
    head = [c.value for c in ws[1]]
    col = head.index("Состав") + 1
    name_col = head.index("Название") + 1
    cells = {
        ws.cell(row=r, column=name_col).value: ws.cell(row=r, column=col).value
        for r in range(2, ws.max_row + 1)
    }
    assert cells["Расходники, реактивы, посуда"] == (
        "Перчатки; Бутыли; Стаканы …ещё 2"
    )
    assert cells["Компаунд заливочный 20 т"] in (None, ""), "без состава ячейка пустая"


# ── Оформление листов (перенесено из тестов старого excel_exporter) ──────────


def test_ссылка_синяя_подчёркнутая_и_текстом_сам_адрес(db):
    """Ячейка «Ссылка» — сам адрес, а не «открыть»: его копируют в письма."""
    _add(db, "1")
    ws = _book(db)["Все активные"]
    head = [c.value for c in ws[1]]
    cell = ws.cell(row=2, column=head.index("Ссылка") + 1)
    assert cell.value == "https://example.test/tender-1/"
    assert cell.hyperlink.target == cell.value
    assert cell.font.color.rgb.endswith("0563C1")
    assert cell.font.underline == "single"


def test_автофильтр_на_каждом_листе_лотов_ровно_по_данным(db):
    _add(db, "1")
    _add(db, "2", name="Компаунд заливочный 30 т")
    wb = _book(db)
    for ws in wb.worksheets:
        if ws.title == "Сводка по темам":
            continue
        rows = ws.max_row - 1
        if rows:
            assert ws.auto_filter.ref == f"A1:Q{rows + 1}", ws.title
        else:
            # Пустой лист: фильтра нет, и книга от этого не ломается.
            assert ws.auto_filter.ref is None, ws.title


def test_площадка_и_номер_процедуры_в_строке(db):
    _add(db, "1", platform="Tender.pro", lot_key="Tender.pro|1", tender_no="TP-77")
    ws = _book(db)["Все активные"]
    assert _names(ws, "Площадка") == ["Tender.pro"]
    assert _names(ws, "№ процедуры") == ["TP-77"]
