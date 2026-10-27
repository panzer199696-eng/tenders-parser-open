"""
Excel-отчёт по темам из единой базы (ТЗ-5, задача 3).

Единственная книга программы (прежний `excel_exporter`, строивший книгу из
лотов текущего прогона по строительному классификатору, удалён 16.10.2026:
его никто не вызывал). Модуль читает готовую базу `tenderwork.db` —
книга собирается всегда, когда база есть, даже если парсинг в этой сессии
не запускался: ленты «что нашлось вчера» не должно пропадать только потому,
что сегодня программу только открыли.

Принципы разметки:

* **Молчание запрещено.** В сводке строка есть у каждой включённой группы
  справочника, даже с нулём лотов: пустая тема — сигнал пересмотреть фразы
  поиска, а не повод спрятать строку.
* **Отсеянное видно поимённо.** Лист «Не наш профиль» присутствует всегда:
  если отфильтрованное не видно, потерю нельзя заметить.
* Оформление прежней книги сохранено (шапка, закрепление, фильтр,
  границы), чтобы человек не переучивался.
"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..dates import parse_deadline, resolve_now
from ..db.feed import classify_lots
from ..store.catalog import load_nomenclature
from ..watch_clients import is_watched, watch_names

logger = logging.getLogger(__name__)

# ── Оформление: как в прежнем экспортёре, чтобы книга выглядела знакомой ─────

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
# Красный для «осталось ≤ 3 дней» — тот же сигнал, что и 🔥 в ленте.
DAYS_DANGER_FONT = Font(color="DC2626", bold=True)
LINK_FONT = Font(color="0563C1", underline="single")


# Пороги цвета оценки — по ТЗ-5 (60/40/20), не как в старом экспортёре:
# темы справочника дают меньшие веса, и старые пороги красили почти всё
# в один серый.
def _score_color(score: int) -> str:
    if score >= 60:
        return "059669"
    if score >= 40:
        return "2563EB"
    if score >= 20:
        return "D97706"
    return "475569"


CLIENT_LABELS = {"full": "полная работа", "partial": "совместно", "market": "рынок"}
INDUSTRY_LABELS = {
    "client": "наш клиент",
    "ours": "отрасль наша",
    "unknown": "неизвестна",
    "foreign": "чужая",
}
# Третье состояние предоплаты нельзя выдавать за «нет»: «не нашли в тексте»
# и «предоплаты нет» — разные выводы для кассового планирования.
PREPAYMENT_LABELS = {"yes": "есть", "no": "нет", "unknown": "не найдено"}

# Колонки листов с лотами — порядок по ТЗ-5.
HEADERS = (
    "№",
    "Оценка",
    "Почему",
    "Темы",
    "Заказчик",
    "Клиент",
    "Отрасль",
    "Площадка",
    "№ процедуры",
    "Название",
    "Состав",
    "Опубликован",
    "Дедлайн",
    "Осталось дней",
    "Предоплата",
    "Сумма, ₽",
    "Ссылка",
)
# Перенос текста нужен там, где строки длинные: без него лист «в ширину»
# не читается, а высота строк скачет.
WRAP_COLUMNS = {3, 4, 10, 11}  # Почему, Темы, Название, Состав
COL_WIDTHS = (5, 8, 46, 24, 24, 14, 13, 12, 12, 60, 50, 12, 12, 9, 11, 14, 40)

FORBIDDEN_SHEET_CHARS = set("[]:*?/\\")
MAX_SHEET_NAME = 31


def _sheet_title(title: str, used: set[str]) -> str:
    """Имя листа из названия темы: ≤31 символа, без запрещённых, уникальное.

    Длинное название режется по границе слова: «Гидроциклоны и пульповое
    хозяйство» должно стать «Гидроциклоны и пульповое», а не
    «Гидроциклоны и пульповое хозяйс». openpyxl молча не режет имена, а
    недопустимое — роняет сохранение книги.
    """
    cleaned = "".join(
        " " if ch in FORBIDDEN_SHEET_CHARS else ch for ch in str(title or "")
    )
    cleaned = " ".join(cleaned.split())
    if len(cleaned) > MAX_SHEET_NAME:
        head = cleaned[:MAX_SHEET_NAME]
        # Если граница слова есть — режем по ней; сплошная длинная строка
        # режется как есть, других вариантов нет.
        if " " in head:
            head = head.rsplit(" ", 1)[0]
        cleaned = head.strip()
    base = cleaned or "Тема"
    candidate = base
    n = 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[: MAX_SHEET_NAME - len(suffix)].rstrip() + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _get(row, key: str, default=None):
    """Значение колонки или default, если колонки в старой базе нет.

    build_report обязан работать и на базе до миграции (книгу могут
    попросить построить до первого ensure_schema) — чтение не должно
    требовать записи.
    """
    try:
        value = row[key]
    except (IndexError, KeyError):
        return default
    return default if value is None else value


def _items_cell(items: list[dict]) -> str:
    """Первые три позиции через «; » и «…ещё N», если позиций больше.

    Количество и единицу в ячейку не пишем: она и так длинная, а состав
    нужен, чтобы глазами решить «наше/не наше» — счёт позиций этому
    не помогает.
    """
    if not items:
        return ""
    names = [str(it.get("name") or "").strip() for it in items[:3]]
    cell = "; ".join(n for n in names if n)
    if len(items) > 3:
        more = f"…ещё {len(items) - 3}"
        cell = f"{cell} {more}" if cell else more
    return cell


def _load_active_lots(conn, today: date, now: datetime) -> list[dict]:
    """Все неистёкшие неразобранные лоты с корзиной, темами и остатком дней.

    Раскладка — та же функция, что у ленты (feed.classify_lots): лист книги
    и корзина ленты на одной базе обязаны совпадать (Г4, 16.10.2026).
    Взятые в работу и отклонённые («Мимо») в книгу не идут — как в ленте.
    Лот без распознанного дедлайна считаем живым: «Не указана» не значит
    «прошёл» — такой лот человек должен увидеть и решить сам.
    """
    try:
        rows = [
            lot for lot in classify_lots(conn, now, today) if lot["bucket"] != "expired"
        ]
    except Exception:
        logger.exception("Не удалось прочитать лоты из базы")
        rows = []

    # Состав закупки — одним запросом на всю книгу (ТЗ-10): точечные
    # заходы по каждому лоту превратились бы в тысячи обращений к базе.
    # Книгу могут собирать и на базе до миграции, где таблицы lot_items
    # ещё нет, — тогда состава просто нет, а не ошибка.
    try:
        from ..db.items import lot_items_by_key

        items_map = lot_items_by_key(conn)
    except Exception:
        items_map = {}

    lots: list[dict] = []
    for row in rows:
        lots.append(
            {
                "row": row,
                "tags": [str(t) for t in row.get("tags_list") or []],
                "lot_key": _get(row, "lot_key", ""),
                "bucket": row["bucket"],
                "score": int(_get(row, "score", 0) or 0),
                "left": row.get("days_left"),
                "items": items_map.get(_get(row, "lot_key", ""), []),
                # Отсеянное профилем: чужая отрасль либо подряд/услуга.
                # Считаем здесь, чтобы ни один лист не забыл про это:
                # «Установка насосной станции. Выполнение
                # строительно-монтажных работ» иначе ложится на лист
                # «Насосы» — ровно та жалоба, с которой начинали.
                "rejected": bool(
                    _get(row, "off_industry", 0) or _get(row, "off_profile", 0)
                ),
            }
        )
    return lots


def _when_collected(conn, run_id: int | None = None) -> str:
    """Строка «когда собрано»: из таблицы runs, иначе из last_run.json.

    База — источник истины (ТЗ-5, задача 5), но в ней может ещё не быть
    ни одной записи: файловый журнал старше и переживает обновление.
    `run_id` — прогон, ради которого строится книга: без него шапка ночной
    книги показывала ПРЕДЫДУЩИЙ прогон (16.10.2026: «15.10 в 07:06»).
    """
    from .. import run_journal
    from ..db.runs import run_by_id

    record = run_by_id(conn, run_id)
    if record is None:
        record = run_journal.read_last_run() or {}
    raw = record.get("finished_at") or record.get("started_at") or ""
    try:
        dt = datetime.fromisoformat(str(raw))
    except ValueError:
        return "неизвестно когда"
    tail = " (догоняющий прогон)" if record.get("trigger") == "catchup" else ""
    return f"{dt:%d.%m} в {dt:%H:%M}{tail}"


def _build_platforms_block(ws, conn, start_row: int, run_id: int | None = None) -> int:
    """Строка на площадку: сколько объектов обошли и всё ли собрали.

    Возвращает номер следующей свободной строки. Если журнала ещё нет
    (первый запуск, старая база), блок не рисуется вовсе — выдумывать
    «всё в порядке» нельзя. Для догоняющего прогона — ночь, дополненная
    догнанными площадками (src.db.runs.run_platforms_for).
    """
    from ..db.runs import run_platforms_for

    rows = run_platforms_for(conn, run_id)
    if not rows:
        return start_row

    title = ws.cell(start_row, 1, "Полнота сбора по площадкам:")
    title.font = Font(bold=True, size=11)
    row_no = start_row + 1
    for r in rows:
        complete = r.get("complete")
        if complete is None:
            verdict = "площадка не сообщила объём"
            color = "808080"
        elif complete:
            # Формулировку берём из журнала: там уже различили «собрано
            # целиком» и «площадка не сообщает объём выдачи».
            verdict = r.get("reason") or "собрано целиком"
            color = "006100"
        else:
            verdict = r.get("reason") or "собрано не всё"
            color = "9C0006"
        ws.cell(row_no, 1, f"   {r.get('platform')}")
        ws.cell(row_no, 2, r.get("collected"))
        cell = ws.cell(row_no, 3, verdict)
        cell.font = Font(color=color, bold=not complete)
        row_no += 1
    return row_no + 1


# Корзины, из которых собираются листы тем. Тема — срез по предмету
# закупки, а не корзина: лот виден и на листе своей корзины, и на листе
# темы. «На проверку» входит: уверенно распознанный «второй эшелон» рынка
# (футеровки, ГШО у незнакомого заказчика) лента копит именно там, а в
# книге его ищут по листу темы — без этого половина листов тем пропала бы
# (сверка на живой базе 16.10.2026). Снятое площадкой и чужое — не входит.
TOPIC_BUCKETS = ("urgent", "core", "client", "below", "review")


def _by_topic(lots: list[dict]) -> dict[str, list[dict]]:
    """Лоты листов тем: по сильнейшей (первой) теме лота.

    На лист темы идёт только уверенно распознанное. Лот, у которого
    совпало общее слово («насосное оборудование», «датчики»), тему лишь
    предполагает — ему место на листе «На проверку».
    Отсеянное профилем не идёт и тогда, когда заказчик на особом контроле
    держит его в рабочем потоке: «Насосная станция. СМР» на листе
    «Насосы» — ровно та жалоба, с которой начинали.
    """
    by_topic: dict[str, list[dict]] = {}
    for lot in lots:
        if (
            lot["bucket"] in TOPIC_BUCKETS
            and lot["tags"]
            and not lot["rejected"]
            and not _get(lot["row"], "needs_review", 0)
        ):
            by_topic.setdefault(lot["tags"][0], []).append(lot)
    return by_topic


def _build_summary_sheet(
    ws, groups: list[dict], lots: list[dict], conn, run_id: int | None = None
) -> None:
    """Сводка по темам: строка на каждую включённую группу, даже с нулём."""
    ws.title = "Сводка по темам"
    # Сверху — происхождение данных: без него книга «в вакууме» и доверия
    # к нулям нет (а вдруг это вчерашняя?)
    info = ws.cell(1, 1, f"Когда собрано: {_when_collected(conn, run_id)}")
    info.font = Font(bold=True, size=11)
    total = ws.cell(2, 1, f"Активных лотов: {len(lots)}")
    total.font = Font(bold=True, size=11, color="1F4E79")

    # Полнота сбора по площадкам (ТЗ-8). Без неё книга молчит о главном: не
    # «сколько нашли», а «всё ли смогли посмотреть». Площадка, отдавшая
    # половину, выглядит в отчёте так же, как отдавшая всё, — и потерю
    # заметить нельзя.
    next_row = _build_platforms_block(ws, conn, start_row=3, run_id=run_id)

    headers = (
        "Тема",
        "Вес",
        "Режим прогона",
        "Лотов",
        "Горит (≤3 дн.)",
        "Ядро профиля",
        "Лучшая оценка",
    )
    header_row = max(next_row, 4)
    for col, title in enumerate(headers, 1):
        cell = ws.cell(header_row, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[header_row].height = 28

    # Лоты темы — ровно те, что лягут на листы тем (_by_topic): иначе
    # сводка и листы показывали бы разные числа.
    by_topic = _by_topic(lots)

    r = header_row
    for group in groups:
        topic = by_topic.get(group["id"], [])
        burning = sum(1 for lot in topic if lot["bucket"] == "urgent")
        core = sum(1 for lot in topic if _get(lot["row"], "core_profile", 0))
        best = max((lot["score"] for lot in topic), default=0)
        r += 1
        values = (
            group["title"],
            group["weight"],
            group["tier"],
            len(topic),
            burning,
            core,
            best if topic else None,
        )
        for col, value in enumerate(values, 1):
            cell = ws.cell(r, col, value)
            cell.border = BORDER
            if col == 7 and value is not None:
                cell.font = Font(color=_score_color(best), bold=True)
                cell.alignment = Alignment(horizontal="center")
            elif col >= 2:
                cell.alignment = Alignment(horizontal="center")
            if not topic and col == 1:
                # Нулевая тема — сигнал: выделяем её молчание, а не прячем.
                cell.font = Font(color="94A3B8", italic=True)

    for col_letter, width in zip("ABCDEFG", (44, 7, 14, 9, 13, 13, 14)):
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    if r > header_row:
        ws.auto_filter.ref = f"A{header_row}:G{r}"


def _date_cell(value) -> tuple[object, str | None]:
    """Дата для ячейки и её формат; непонятое остаётся текстом как есть."""
    at = parse_deadline(value)
    if at is None:
        return value or "", None
    if re.search(r"\d{1,2}:\d{2}", str(value)):
        return at, "DD.MM.YYYY HH:MM"
    return datetime(at.year, at.month, at.day), "DD.MM.YYYY"


def _link_cell(row) -> str:
    """Ссылка на первоисточник, а без ссылки — номер, по которому искать.

    Пустая ячейка у лота без ссылки (у ТЭК-Торга и Bidzaar таких сотни)
    оставляла человека без пути к лоту; номер процедуры ищется на площадке.
    """
    href = _get(row, "main_href", None)
    if href:
        return href
    number = _get(row, "tender_no", "") or _get(row, "lot_id", "")
    return f"№ {number} — искать на площадке" if number else ""


def _build_lots_sheet(ws, lots: list[dict], titles: dict[str, str]) -> None:
    """Лист с лотами: колонки и оформление по образцу прежней книги."""
    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[1].height = 28

    for idx, lot in enumerate(lots, 1):
        row = lot["row"]
        published, published_format = _date_cell(_get(row, "date_published", ""))
        deadline, deadline_format = _date_cell(_get(row, "date_deadline", ""))
        link = _link_cell(row)
        values = (
            idx,
            lot["score"],
            _get(row, "score_explain", "") or "",
            ", ".join(titles.get(t, t) for t in lot["tags"]) or "—",
            _get(row, "company", "") or "",
            CLIENT_LABELS.get(_get(row, "client_level", ""), "—"),
            INDUSTRY_LABELS.get(_get(row, "industry", "unknown"), "неизвестна"),
            _get(row, "platform", "") or "",
            _get(row, "tender_no", "") or "",
            _get(row, "name", "") or "",
            _items_cell(lot.get("items")),
            published,
            deadline,
            lot["left"] if lot["left"] is not None else "—",
            PREPAYMENT_LABELS.get(_get(row, "prepayment", "unknown"), "не найдено"),
            _get(row, "sum", None),
            link,
        )
        r = idx + 1
        for col, value in enumerate(values, 1):
            cell = ws.cell(r, col, value)
            cell.border = BORDER
            if col in WRAP_COLUMNS:
                cell.alignment = Alignment(
                    horizontal="left", vertical="top", wrap_text=True
                )

        # Даты (L=12, M=13) — ячейки-даты: фильтр и сортировка по сроку в
        # книге работают только с датой, не с текстом (Г3, 16.10.2026).
        if published_format:
            ws.cell(r, 12).number_format = published_format
        if deadline_format:
            ws.cell(r, 13).number_format = deadline_format

        # Оценка — цветной жирный шрифт (B=2), как в прежней книге.
        score_cell = ws.cell(r, 2)
        score_cell.font = Font(color=_score_color(lot["score"]), bold=True)
        score_cell.alignment = Alignment(horizontal="center")

        # «Осталось дней» (N=14): красным, когда до дедлайна ≤ 3 дней.
        days_cell = ws.cell(r, 14)
        days_cell.alignment = Alignment(horizontal="center")
        if lot["left"] is not None and lot["left"] <= 3:
            days_cell.font = DAYS_DANGER_FONT

        # Сумма (P=16) — число с разделителем разрядов: глазу нужен масштаб.
        sum_value = _get(row, "sum", None)
        if sum_value is not None:
            ws.cell(r, 16).number_format = "#,##0"
            ws.cell(r, 16).alignment = Alignment(horizontal="right")

        # Ссылка (Q=17) — кликабельная; «№ … — искать на площадке» — текст.
        link_cell = ws.cell(r, 17)
        if _get(row, "main_href", None):
            link_cell.hyperlink = link_cell.value
            link_cell.font = LINK_FONT
            link_cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A2"
    if lots:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(lots) + 1}"


def _watch_sheet_title(names: list[str]) -> str:
    """Название листа: пока заказчиков двое-трое, называем их поимённо.

    «Ромашка и Лютик» человеку понятнее, чем «Особый контроль»: видно, что
    внутри, не открывая лист.
    """
    if len(names) <= 3:
        head = ", ".join(names[:-1])
        return f"{head} и {names[-1]}" if len(names) > 1 else names[0]
    return "На особом контроле"


def _is_watched(lot: dict, names: list[str]) -> bool:
    """Лот этого заказчика? Тонкая обёртка над watch_clients.is_watched:
    берёт организатора из строки базы через _get, чтобы не переписывать
    вызовы в build_report. Само правило (сверка по целым словам) живёт
    в одном месте — src/watch_clients.py.
    """
    return is_watched(str(_get(lot["row"], "company", "") or ""), names)


def build_report(
    conn,
    today: date | None = None,
    now: datetime | None = None,
    nomenclature: dict | None = None,
    run_id: int | None = None,
) -> bytes:
    """Собрать книгу Excel из единой базы. Возвращает байты .xlsx.

    `today`/`now` задаются в тестах: срочность и истечение зависят от
    текущего дня, и проверять их можно только на фиксированной дате.
    `run_id` — прогон для шапки и блока полноты (None — последний в базе).
    """
    today, now = resolve_now(today, now)
    nomenclature = nomenclature if nomenclature is not None else load_nomenclature()

    # Только включённые группы: выключенная тема не ищется и не считается —
    # строка про неё была бы шумом.
    groups = sorted(
        (
            {
                "id": str(g["id"]),
                "title": str(g.get("title") or g["id"]),
                "weight": int(g.get("weight", 0)),
                "tier": str(g.get("tier", "off")),
            }
            for g in nomenclature.get("groups", [])
            if g.get("id") and g.get("enabled", True)
        ),
        key=lambda g: -g["weight"],
    )
    titles = {g["id"]: g["title"] for g in groups}

    lots = _load_active_lots(conn, today, now)
    wb = Workbook()

    _build_summary_sheet(wb.active, groups, lots, conn, run_id)

    # Листы корзин строятся по той же раскладке, что лента (feed.classify_row):
    # лист и корзина на одной базе совпадают (Г4, 16.10.2026).
    def bucket(key: str) -> list[dict]:
        return [lot for lot in lots if lot["bucket"] == key]

    # Горит — по правилу ленты (16.10.2026): наше, закреплённые
    # клиенты и особый контроль со сроком ≤ 3 дней. Раньше книга жгла любой
    # лот с оценкой > 0, и «рынок» с дедлайном завтра стоял рядом с нашим.
    # Сортировка по сроку: из двух горящих раньше разбирается тот, у кого
    # меньше времени.
    burning = sorted(bucket("urgent"), key=lambda lot: (lot["left"], -lot["score"]))
    _build_lots_sheet(wb.create_sheet("🔥 Горит"), burning, titles)

    # Лист на каждую тему: только непустые, порядок по весу группы —
    # верх книги отвечает на «что для нас важнее».
    used_names: set[str] = set()
    by_topic = _by_topic(lots)
    for group in groups:
        topic = by_topic.get(group["id"])
        if not topic:
            continue
        topic.sort(key=lambda lot: -lot["score"])
        ws = wb.create_sheet(_sheet_title(group["title"], used_names))
        _build_lots_sheet(ws, topic, titles)

    # Закреплённые клиенты: за заказчиком, с которым мы работаем, смотрим
    # весь поток целиком — даже то, чью тему справочник не распознал. Наша
    # позиция часто спрятана внутри сборного лота («Расходники, реактивы,
    # посуда»), где ни одна фраза не сработает. Здесь же — лот особого
    # контроля, отсеянный профилем: как в ленте, а не в «Не наш профиль».
    clients = sorted(bucket("client"), key=lambda lot: -lot["score"])
    _build_lots_sheet(wb.create_sheet("Закреплённые клиенты"), clients, titles)

    # Заказчики на особом контроле — отдельным листом, до общих корзин:
    # по ним работа идёт адресно, и искать их среди тысячи строк незачем.
    # Сами имена и сверка — в src/watch_clients.py, здесь только лист.
    watched_names = watch_names()
    if watched_names:
        watched = sorted(
            (lot for lot in lots if _is_watched(lot, watched_names)),
            key=lambda lot: (
                -lot["score"],
                parse_deadline(_get(lot["row"], "date_deadline", "")) or datetime.max,
            ),
        )
        _build_lots_sheet(
            wb.create_sheet(
                _sheet_title(_watch_sheet_title(watched_names), used_names)
            ),
            watched,
            titles,
        )

    # Ниже порога, На проверку, Не наш профиль, Пропал — листы обязательны
    # (даже пустые): отсеянное и сомнительное должно быть видно поимённо,
    # иначе потерю нельзя заметить.
    for key, title in (
        ("below", "Ниже порога"),
        ("review", "На проверку"),
        ("foreign", "Не наш профиль"),
        ("missing", "Пропал — проверить"),
    ):
        rows = sorted(bucket(key), key=lambda lot: -lot["score"])
        _build_lots_sheet(wb.create_sheet(title), rows, titles)

    # Все активные — по убыванию оценки: универсальный вид «на всё».
    active = sorted(lots, key=lambda lot: -lot["score"])
    _build_lots_sheet(wb.create_sheet("Все активные"), active, titles)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
